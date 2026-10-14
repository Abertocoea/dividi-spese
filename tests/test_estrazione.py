from typing import Any

import pymupdf
from fastapi.testclient import TestClient

from conftest import FIXTURES, EstrattoreFinto, output_grezzo
from dividi_spese.estrazione import ErroreOllama, ModelloNonInstallato, OllamaNonRaggiungibile


def pdf_con_testo(testo: str) -> bytes:
    doc = pymupdf.open()
    doc.new_page().insert_text((50, 72), testo)
    return bytes(doc.tobytes())


def pdf_scansionato() -> bytes:
    """PDF senza testo: una sola immagine per pagina, come uno scanner."""
    immagine = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 20), False)
    immagine.clear_with(200)
    doc = pymupdf.open()
    doc.new_page().insert_image(pymupdf.Rect(0, 0, 200, 100), pixmap=immagine)
    return bytes(doc.tobytes())


def png() -> bytes:
    immagine = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 10, 10), False)
    immagine.clear_with(255)
    return bytes(immagine.tobytes("png"))


def estrai(client: TestClient, nome: str, dati: bytes, tipo: str, **form: str) -> Any:
    return client.post("/estrai", files={"file": (nome, dati, tipo)}, data=form)


def test_un_pdf_con_testo_viene_inviato_al_modello_come_testo(client: TestClient, estrattore: EstrattoreFinto) -> None:
    r = estrai(client, "ordine.pdf", pdf_con_testo("1 Silenzio 31 IT NM M10 R 6,00 EUR"), "application/pdf")
    assert r.status_code == 200
    assert "1 Silenzio 31 IT NM M10 R 6,00 EUR" in (estrattore.ricevuti[0].testo or "")
    assert estrattore.ricevuti[0].immagini == []
    prodotti = r.json()["ordine"]["prodotti"]
    assert [(p["nome"], p["quantita"], p["prezzo_unitario"]) for p in prodotti] == [
        ("Silenzio", 1, "6.00"),
        ("Negazione Ostinata", 3, "2.50"),
        ("Laghetto Bollente", 2, "27.00"),
    ]


def test_un_pdf_scansionato_viene_inviato_come_immagini_delle_pagine(client: TestClient, estrattore: EstrattoreFinto) -> None:
    r = estrai(client, "scansione.pdf", pdf_scansionato(), "application/pdf")
    assert r.status_code == 200
    ricevuto = estrattore.ricevuti[0]
    assert ricevuto.testo is None
    assert len(ricevuto.immagini) == 1 and ricevuto.immagini[0].startswith(b"\x89PNG")


def test_uno_screenshot_viene_inviato_come_immagine(client: TestClient, estrattore: EstrattoreFinto) -> None:
    dati = png()
    r = estrai(client, "carrello.PNG", dati, "image/png")
    assert r.status_code == 200
    assert estrattore.ricevuti[0].immagini == [dati]


def test_accetta_screenshot_jpg_e_webp(client: TestClient, estrattore: EstrattoreFinto) -> None:
    for nome in ("carrello.jpg", "carrello.jpeg", "carrello.webp"):
        assert estrai(client, nome, b"immagine", "image/*").status_code == 200
    assert len(estrattore.ricevuti) == 3


def test_rifiuta_un_formato_non_supportato(client: TestClient, estrattore: EstrattoreFinto) -> None:
    r = estrai(client, "ordine.docx", b"...", "application/octet-stream")
    assert r.status_code == 415
    assert "PDF, PNG, JPG, WebP, CSV, XLSX" in r.json()["detail"]
    assert estrattore.ricevuti == []


def test_segnala_un_pdf_illeggibile(client: TestClient) -> None:
    r = estrai(client, "ordine.pdf", b"non sono un pdf", "application/pdf")
    assert r.status_code == 422
    assert "illeggibile" in r.json()["detail"]


def test_un_csv_italiano_viene_inviato_come_tabella_testuale(client: TestClient, estrattore: EstrattoreFinto) -> None:
    r = estrai(client, "ordine_it.csv", (FIXTURES / "ordine_it.csv").read_bytes(), "text/csv")
    assert r.status_code == 200
    testo = estrattore.ricevuti[0].testo or ""
    assert "Quantità | Prodotto | Dettagli | Prezzo unitario (EUR)" in testo
    assert "3 | Negazione Ostinata | #56 EN NM KTK U | 2,50" in testo
    assert "Totale |  | 182,00" in testo


def test_un_csv_con_la_virgola_come_separatore_viene_letto_per_colonne(client: TestClient, estrattore: EstrattoreFinto) -> None:
    estrai(client, "ordine_en.csv", (FIXTURES / "ordine_en.csv").read_bytes(), "text/csv")
    assert "Negazione Ostinata | #56 EN NM KTK U | 3 | 7.50" in (estrattore.ricevuti[0].testo or "")


def test_un_csv_in_codifica_windows_viene_letto(client: TestClient, estrattore: EstrattoreFinto) -> None:
    dati = "Quantità;Prodotto;Prezzo\n1;Carta più rara;5,00\n".encode("cp1252")
    assert estrai(client, "ordine.csv", dati, "text/csv").status_code == 200
    assert "1 | Carta più rara | 5,00" in (estrattore.ricevuti[0].testo or "")


def test_un_excel_viene_letto_con_tutti_i_fogli(client: TestClient, estrattore: EstrattoreFinto) -> None:
    r = estrai(client, "ordine.xlsx", (FIXTURES / "ordine.xlsx").read_bytes(), "application/vnd.ms-excel")
    assert r.status_code == 200
    testo = estrattore.ricevuti[0].testo or ""
    assert "## Sheet: Prodotti" in testo and "## Sheet: Riepilogo" in testo
    assert "Spese di spedizione | 11,95 €" in testo
    assert "3 | Negazione Ostinata | #56 EN NM KTK U | 2,50 € | 7,50 €" in testo


def test_l_anteprima_di_csv_ed_excel_contiene_la_tabella_originale(client: TestClient) -> None:
    r = estrai(client, "ordine.xlsx", (FIXTURES / "ordine.xlsx").read_bytes(), "application/vnd.ms-excel")
    fogli = r.json()["anteprima"]
    assert [f["nome"] for f in fogli] == ["Prodotti", "Riepilogo"]
    assert fogli[1]["righe"][0] == ["Voce", "Importo"]
    assert len(fogli[0]["righe"]) == 18


def test_l_anteprima_e_vuota_per_pdf_e_immagini(client: TestClient) -> None:
    r = estrai(client, "carrello.png", png(), "image/png")
    assert r.json()["anteprima"] == []


def test_segnala_un_excel_illeggibile(client: TestClient) -> None:
    r = estrai(client, "ordine.xlsx", b"non sono un excel", "application/vnd.ms-excel")
    assert r.status_code == 422
    assert "illeggibile" in r.json()["detail"]


def test_segnala_che_ollama_non_e_avviato(client: TestClient, estrattore: EstrattoreFinto) -> None:
    estrattore.errore = OllamaNonRaggiungibile()
    r = estrai(client, "carrello.png", png(), "image/png")
    assert r.status_code == 503
    assert "Ollama non è raggiungibile" in r.json()["detail"]


def test_segnala_il_modello_mancante_con_il_comando_per_installarlo(client: TestClient, estrattore: EstrattoreFinto) -> None:
    estrattore.errore = ModelloNonInstallato("qwen3.5:9b-q4_K_M")
    r = estrai(client, "carrello.png", png(), "image/png")
    assert r.status_code == 424
    assert "ollama pull qwen3.5:9b-q4_K_M" in r.json()["detail"]


def test_usa_il_modello_scelto_dall_utente(client: TestClient, estrattore: EstrattoreFinto) -> None:
    estrai(client, "carrello.png", png(), "image/png", modello="qwen3.5:4b-q4_K_M")
    assert estrattore.modelli_usati == ["qwen3.5:4b-q4_K_M"]


def test_riporta_gli_avvisi_dell_estrattore(client: TestClient, estrattore: EstrattoreFinto) -> None:
    estrattore.avvisi = ["Il modello non è tutto sulla GPU: l'estrazione sarà lenta"]
    r = estrai(client, "carrello.png", png(), "image/png")
    assert r.json()["avvisi"] == ["Il modello non è tutto sulla GPU: l'estrazione sarà lenta"]


def test_elenca_i_modelli_con_visione_installati(client: TestClient) -> None:
    r = client.get("/modelli")
    assert r.json() == {"modelli": ["qwen3.5:9b-q4_K_M", "qwen3.5:4b-q4_K_M"], "predefinito": "qwen3.5:9b-q4_K_M"}


def test_l_elenco_modelli_segnala_che_ollama_non_e_avviato(client: TestClient, estrattore: EstrattoreFinto) -> None:
    estrattore.errore = OllamaNonRaggiungibile()
    assert client.get("/modelli").status_code == 503


def test_gli_importi_excel_con_il_formato_euro_italiano_restano_corretti(client: TestClient, estrattore: EstrattoreFinto) -> None:
    import io

    from openpyxl import Workbook

    libro = Workbook()
    foglio = libro.active
    assert foglio is not None
    for riga, formato in enumerate(['#,##0.00\\ "€"', "0.00€", "#,##0.00", "0"], start=1):
        foglio.cell(riga, 1, f"formato {riga}")
        foglio.cell(riga, 2, 1234.5).number_format = formato
    dati = io.BytesIO()
    libro.save(dati)
    estrai(client, "ordine.xlsx", dati.getvalue(), "application/vnd.ms-excel")
    testo = estrattore.ricevuti[0].testo or ""
    assert "formato 1 | 1.234,50 €" in testo
    assert "formato 2 | 1234,50 €" in testo
    assert "formato 3 | 1.234,50" in testo
    assert "formato 4 | 1235" in testo


def test_un_errore_di_ollama_durante_l_estrazione_viene_spiegato(client: TestClient, estrattore: EstrattoreFinto) -> None:
    estrattore.errore = ErroreOllama("model requires more system memory than is available")
    r = estrai(client, "carrello.png", png(), "image/png")
    assert r.status_code == 502
    assert "model requires more system memory" in r.json()["detail"]
